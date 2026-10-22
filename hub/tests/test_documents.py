"""Phase 1B acceptance: Refresh Library's 13 requirements.

Covers: new / changed / unchanged / duplicate-avoidance / missing /
unsupported / unreadable files, failed-update content preservation,
OCR-needed reporting, sheet-name preservation, macro safety (xlsm is
never opened), and refresh timestamp recording.
"""
from pathlib import Path

import pytest

from hubcore import db as dbmod
from hubcore import documents


def make_text_pdf(path: Path, text: str) -> None:
    """Write a minimal valid single-page PDF with extractable text."""
    t = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content = f"BT /F1 24 Tf 100 700 Td ({t}) Tj ET".encode("latin-1")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n"
        + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF"
    ).encode()
    path.write_bytes(bytes(out))


def make_blank_pdf(path: Path) -> None:
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=612, height=792)
    with open(path, "wb") as fh:
        w.write(fh)


@pytest.fixture()
def inbox(tmp_path):
    box = tmp_path / "inbox"
    box.mkdir()
    (box / "notes.txt").write_text("Deferred tax arises from temporary differences.", encoding="utf-8")
    (box / "summary.md").write_text("# IAS 2\nInventories are measured at cost.", encoding="utf-8")
    (box / "ledger.csv").write_text("date,desc,amount\n2026-01-01,Sales,1000\n", encoding="utf-8")
    make_text_pdf(box / "standard.pdf", "Objective of IAS 2 Inventories")
    # xlsx with two sheets
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TrialBalance"
    ws.append(["Account", "Debit", "Credit"])
    ws.append(["Cash", 5000, 0])
    ws2 = wb.create_sheet("Notes")
    ws2.append(["Note", "Text"])
    ws2.append([1, "Going concern"])
    wb.save(box / "tb.xlsx")
    # docx
    import docx

    doc = docx.Document()
    doc.add_paragraph("Audit risk is the risk of material misstatement.")
    doc.save(box / "audit.docx")
    # permitted subfolder (spec req. 1)
    sub = box / "sub"
    sub.mkdir()
    (sub / "deep.txt").write_text("Subfolder documents are indexed too.", encoding="utf-8")
    return box


def test_new_files_indexed(conn, inbox):
    report = documents.refresh_library(conn, inbox)
    assert sorted(report["new"]) == sorted(
        ["notes.txt", "summary.md", "ledger.csv", "standard.pdf", "tb.xlsx",
         "audit.docx", "sub/deep.txt"]
    )
    assert report["failed"] == [] and report["unsupported"] == []
    assert report["indexed"] == 7
    chunks = conn.execute("SELECT COUNT(*) FROM doc_chunks").fetchone()[0]
    assert chunks > 7  # multi-page/sheet files produce several chunks
    fts = conn.execute("SELECT COUNT(*) FROM doc_fts").fetchone()[0]
    assert fts == chunks


def test_unchanged_files_skipped(conn, inbox):
    documents.refresh_library(conn, inbox)
    before = conn.execute("SELECT COUNT(*) FROM doc_chunks").fetchone()[0]
    report = documents.refresh_library(conn, inbox)
    assert len(report["unchanged"]) == 7 and report["new"] == []
    after = conn.execute("SELECT COUNT(*) FROM doc_chunks").fetchone()[0]
    assert before == after  # no duplicate records


def test_modified_file_reindexed_with_stable_id(conn, inbox):
    documents.refresh_library(conn, inbox)
    old_id = conn.execute(
        "SELECT id FROM documents WHERE filename='notes.txt'"
    ).fetchone()[0]
    (inbox / "notes.txt").write_text("Deferred tax arises from TEMPORARY differences updated.", encoding="utf-8")
    report = documents.refresh_library(conn, inbox)
    assert report["updated"] == ["notes.txt"]
    new_id = conn.execute("SELECT id FROM documents WHERE filename='notes.txt'").fetchone()[0]
    assert new_id == old_id  # identity preserved
    hits = documents.search_chunks(conn, "updated")
    assert any("notes.txt" in h["filename"] for h in hits)
    assert documents.search_chunks(conn, "nonexistentwordxyz") == []


def test_missing_file_marked_not_deleted(conn, inbox):
    documents.refresh_library(conn, inbox)
    (inbox / "notes.txt").unlink()
    report = documents.refresh_library(conn, inbox)
    assert any("notes.txt" in m for m in report["missing"])
    row = conn.execute("SELECT * FROM documents WHERE filename='notes.txt'").fetchone()
    assert row["status"] == "missing"  # row kept, never auto-deleted
    # chunks/records preserved for history
    assert conn.execute(
        "SELECT COUNT(*) FROM doc_chunks WHERE document_id=?", (row["id"],)
    ).fetchone()[0] > 0


def test_unsupported_reported_and_never_opened(conn, inbox):
    (inbox / "macro.xlsm").write_bytes(b"MZ fake executable")
    (inbox / "run.exe").write_bytes(b"MZ fake")
    report = documents.refresh_library(conn, inbox)
    assert sorted(report["unsupported"]) == ["macro.xlsm", "run.exe"]
    assert report["failed"] == []
    assert conn.execute("SELECT COUNT(*) FROM documents").fetchone()[0] == 7


def test_corrupt_pdf_failed_but_refresh_continues(conn, inbox):
    (inbox / "broken.pdf").write_bytes(b"%PDF-1.4 this is not a pdf")
    report = documents.refresh_library(conn, inbox)
    assert report["failed"] == ["broken.pdf"]
    assert len(report["new"]) == 7  # everything else still indexed
    row = conn.execute("SELECT * FROM documents WHERE filename='broken.pdf'").fetchone()
    assert row["status"] == "failed" and row["error_text"]


def test_scanned_pdf_reported_ocr_needed(conn, inbox):
    make_blank_pdf(inbox / "scanned.pdf")
    report = documents.refresh_library(conn, inbox)
    assert report["ocr_needed"] == ["scanned.pdf"]
    row = conn.execute("SELECT * FROM documents WHERE filename='scanned.pdf'").fetchone()
    assert row["status"] == "ocr_needed"


def test_xlsx_sheet_names_preserved(conn, inbox):
    documents.refresh_library(conn, inbox)
    sections = {
        r[0]
        for r in conn.execute(
            "SELECT DISTINCT sheet_name FROM doc_chunks "
            "WHERE document_id=(SELECT id FROM documents WHERE filename='tb.xlsx')"
        )
    }
    assert sections == {"TrialBalance", "Notes"}
    hits = documents.search_chunks(conn, "TrialBalance")
    assert hits  # sheet name is searchable


def test_failed_update_preserves_previous_content(conn, inbox):
    (inbox / "fragile.pdf").write_bytes(b"placeholder")
    make_text_pdf(inbox / "fragile.pdf", "Original valid content here")
    documents.refresh_library(conn, inbox)
    assert documents.search_chunks(conn, "Original valid content")
    # now corrupt it: hash changes, extraction fails
    (inbox / "fragile.pdf").write_bytes(b"%PDF-1.4 garbage garbage")
    report = documents.refresh_library(conn, inbox)
    assert report["failed"] == ["fragile.pdf"]
    row = conn.execute("SELECT * FROM documents WHERE filename='fragile.pdf'").fetchone()
    assert row["status"] == "indexed"  # previous valid state kept
    assert documents.search_chunks(conn, "Original valid content")  # old chunks kept


def test_refresh_timestamp_recorded(conn, inbox):
    report = documents.refresh_library(conn, inbox)
    assert report["started_at"] and report["finished_at"]
    saved = conn.execute(
        "SELECT value_json FROM settings WHERE key='last_refresh'"
    ).fetchone()[0]
    assert "finished_at" in saved


def test_missing_inbox_dir_does_not_crash(conn, tmp_path):
    report = documents.refresh_library(conn, tmp_path / "nope")
    assert "error" in report
    assert report["new"] == []


def test_search_chunks_snippet(conn, inbox):
    documents.refresh_library(conn, inbox)
    hits = documents.search_chunks(conn, "Inventories")
    assert hits
    by_file = {h["filename"] for h in hits}
    assert "summary.md" in by_file and "standard.pdf" in by_file
    assert all("<b>" in h["snippet"] for h in hits)  # highlighted matches
