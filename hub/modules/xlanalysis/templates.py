"""Python-generated Excel workbooks (openpyxl). Formatting only — no VBA,
no macros, and embedded macros from source files are never executed.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_TITLE_FONT = Font(bold=True, size=14)
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="2F5496")
_FLAG_FILL = PatternFill("solid", fgColor="FCE4EC")
_THIN = Side(style="thin", color="B0B0B0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)


def _autosize(ws, ncols, min_width=12, max_width=40):
    for c in range(1, ncols + 1):
        width = min_width
        for cell in ws[get_column_letter(c)]:
            if cell.value:
                width = max(width, min(len(str(cell.value)) + 2, max_width))
        ws.column_dimensions[get_column_letter(c)].width = width


def variance_report_wb(result: dict, title: str, path: str | Path) -> Path:
    """Formatted variance report: summary + flagged detail."""
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Variance report"
    ws["A1"] = title
    ws["A1"].font = _TITLE_FONT
    ws["A3"] = "Totals"
    for i, (k, v) in enumerate(result["totals"].items(), start=4):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    params = result.get("params", {})
    ws.cell(row=9, column=1,
            value=f"Thresholds: |variance %| >= {params.get('threshold_pct')} "
                  f"and |variance| >= {params.get('threshold_abs')}")
    ws.cell(row=10, column=1,
            value=f"Flagged rows: {result['flagged_count']} of {len(result['rows'])}")

    start = 12
    params = result.get("params", {})
    la = params.get("label_actual", "Actual")
    lc = params.get("label_compare", "Budget")
    headers = ["Row", "Group", la, lc, "Variance", "Variance %",
               "Flagged"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=start, column=c, value=h)
    _style_header(ws, start, len(headers))
    rows = result["rows"]
    for r_i, r in enumerate(rows, start=start + 1):
        vals = [r["row"], r["group"], r[la], r[lc],
                r["variance"],
                r["variance_pct"] * 100 if r["variance_pct"] is not None else None,
                "YES" if r["flagged"] else ""]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r_i, column=c, value=v)
            cell.border = _BORDER
            if r["flagged"]:
                cell.fill = _FLAG_FILL
    _autosize(ws, len(headers))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def summary_wb(df, title: str, path: str | Path) -> Path:
    """Formatted group-by summary workbook."""
    import pandas as pd

    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = title
    ws["A1"].font = _TITLE_FONT
    headers = list(df.columns)
    for c, h in enumerate(headers, start=1):
        ws.cell(row=3, column=c, value=h)
    _style_header(ws, 3, len(headers))
    for r_i, (_, row) in enumerate(df.iterrows(), start=4):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r_i, column=c,
                           value=None if pd.isna(v) else v)
            cell.border = _BORDER
    _autosize(ws, len(headers))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def findings_wb(findings: list[dict], title: str, path: str | Path) -> Path:
    """Formatted exception-findings workbook."""
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Exceptions"
    ws["A1"] = title
    ws["A1"].font = _TITLE_FONT
    headers = ["Rule", "Severity", "Row", "Column", "Value", "Message"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=3, column=c, value=h)
    _style_header(ws, 3, len(headers))
    for r_i, f in enumerate(findings, start=4):
        for c, v in enumerate(
                [f["rule"], f["severity"], f["row"], f["column"],
                 f["value"], f["message"]], start=1):
            cell = ws.cell(row=r_i, column=c, value=v)
            cell.border = _BORDER
            if f["severity"] == "error":
                cell.fill = _FLAG_FILL
    _autosize(ws, len(headers))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
